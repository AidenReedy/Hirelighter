"""Build a Word version of the resume that mirrors the LaTeX layout."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .escape import bold_runs

FONT = "Times New Roman"
TEXT_WIDTH = Inches(7.5)  # letter, 0.5in margins


def _fmt(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return p


def _run(p, text, size=10, bold=False, italic=False, small_caps=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    r.font.small_caps = small_caps
    return r


def _rich(p, text, size=10, italic=False):
    for t, b in bold_runs(text):
        _run(p, t, size=size, bold=b, italic=italic)


def _hyperlink(p, url, text, size=10):
    part = p.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:u", "single"), ("w:sz", str(size * 2))):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), FONT)
            el.set(qn("w:hAnsi"), FONT)
        else:
            el.set(qn("w:val"), val)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    p._p.append(link)


def _bottom_border(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _two_col(doc, left, right, size=10, bold=False, italic=False, indent=0.15):
    p = _fmt(doc.add_paragraph(), 0, 0)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    if isinstance(left, list):
        for t, kw in left:
            _run(p, t, size=size, **kw)
    else:
        _run(p, left.replace("--", "–"), size=size, bold=bold, italic=italic)
    if right:
        _run(p, "\t" + right.replace("--", "–"), size=size if not bold else 10, italic=italic)
    return p


def render_docx(doc: dict) -> bytes:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin = sec.bottom_margin = Inches(0.5)

    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    prof = doc["profile"]
    p = _fmt(d.add_paragraph(), 0, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, prof["name"], size=24, bold=True, small_caps=True)

    p = _fmt(d.add_paragraph(), 0, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, l in enumerate(prof["links"]):
        if i:
            _run(p, "  |  ")
        if l.get("url"):
            _hyperlink(p, l["url"], l["text"])
        else:
            _run(p, l["text"])

    for s in doc["sections"]:
        h = _fmt(d.add_paragraph(), 6, 3)
        _run(h, s["title"], size=12, small_caps=True)
        _bottom_border(h)

        if s["kind"] == "skills":
            sep = None if doc.get("skills_layout") == "lines" else "   |   "
            p = _fmt(d.add_paragraph(), 0, 0)
            p.paragraph_format.left_indent = Inches(0.15)
            for i, g in enumerate(s["groups"]):
                if i:
                    if sep:
                        _run(p, sep)
                    else:
                        p = _fmt(d.add_paragraph(), 0, 0)
                        p.paragraph_format.left_indent = Inches(0.15)
                _run(p, g["name"], bold=True)
                _run(p, ": " + ", ".join(g["skills"]))
            continue

        for e in s["entries"]:
            k = e["kind"]
            if k == "job":
                _two_col(d, e["title"], e["date"], size=11, bold=True).paragraph_format.space_before = Pt(2)
                _two_col(d, e["org"], e["location"], italic=True)
            elif k == "education":
                _two_col(d, e["org"], e["location"], size=11, bold=True).paragraph_format.space_before = Pt(2)
                _two_col(d, e["title"], e["date"], italic=True)
            elif k == "cert":
                _two_col(d, e["title"], e["date"], size=11, bold=True).paragraph_format.space_before = Pt(2)
            else:
                left = [(e["title"], {"bold": True})]
                if e.get("tech"):
                    left += [(" | ", {}), (e["tech"], {"italic": True})]
                _two_col(d, left, e["date"]).paragraph_format.space_before = Pt(2)
            for b in e["bullets"]:
                p = _fmt(d.add_paragraph(style="List Bullet"), 0, 0)
                p.paragraph_format.left_indent = Inches(0.45)
                p.paragraph_format.first_line_indent = Inches(-0.17)
                _rich(p, b)

    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()
